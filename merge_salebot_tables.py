"""Download SaleBot shared tables and combine their CSV exports into one XLSX."""

import argparse
import csv
import json
import sys
from pathlib import Path

import requests
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


SOURCES = [
    ("01_7_materials", "7 материалов", "JHB45bhtrRG1qMYpoExqFDF4h-hncNrYcgvtZf3Z7wI"),
    ("02_ege_readiness", "Тест на готовность к ЕГЭ", "3kR3e9vGKW4WsHsSfLcd3UfljgcEJPVpJKPzoCXuf88"),
    ("03_practice_exams", "Пробные варианты", "Y9jj6-B7eQVc-aPjozwLjLh95zCVxqxxgGFkWiJGhxQ"),
    ("04_admission_calculator", "Калькулятор поступления", "W4pc5dUPavqz4Pi5zfLMngDKmTJtPs5T9UjQggiCfxQ"),
    ("05_plan_generator", "Генератор плана", "J45s5XNX5XTf2Z7SWhF7l11LxGUMJ_KzfkgjNq4ahqM"),
    ("06_ege_theory", "Вся теория для сдачи ЕГЭ", "ZS-P5fAx7kQh-TOVidgZ6puE-enzAbfU4EkEV4t-0lU"),
    ("07_parent_kit", "Набор для родителя", "cPNWh8sBjDli9Y0gR4S8QPRv4BSHgQvg4XwNh0YNikI"),
    ("08_essay_books", "Книги для сочинений", "jGIb_CoS2zXrdYPgZhQIqm1lMgZswALlMKivbkuEVc8"),
    ("09_courses", "Курсы", "LyhO3dX29BqjKwIXIn3XrdAUBTKpyGAQJkgFmOZPzTE"),
]


def download_sources(directory: Path, missing_only: bool = False) -> None:
    directory.mkdir(exist_ok=True)
    with requests.Session() as session:
        for filename, label, token in SOURCES:
            target = directory / f"{filename}.csv"
            if missing_only and target.exists():
                continue
            temporary = target.with_suffix(".csv.part")
            url = f"https://salebot.pro/shared/table/{token}/export.csv"
            print(f"Downloading {label}...", flush=True)
            with session.get(url, stream=True, timeout=(15, 180)) as response:
                response.raise_for_status()
                content_type = response.headers.get("Content-Type", "")
                if "csv" not in content_type:
                    raise ValueError(f"Unexpected content type for {label}: {content_type}")
                with temporary.open("wb") as output:
                    for chunk in response.iter_content(chunk_size=1024 * 1024):
                        if chunk:
                            output.write(chunk)
            temporary.replace(target)
            print(f"  {target.name}: {target.stat().st_size} bytes", flush=True)


def source_reader(path: Path):
    source = path.open("r", encoding="utf-8-sig", newline="")
    return source, csv.reader(source, delimiter=";")


def inspect_sources(directory: Path):
    all_headers = []
    summaries = []
    for filename, label, _ in SOURCES:
        path = directory / f"{filename}.csv"
        source, reader = source_reader(path)
        with source:
            headers = next(reader)
            if len(headers) != len(set(headers)):
                raise ValueError(f"Repeated column names in {path}")
            for header in headers:
                if header not in all_headers:
                    all_headers.append(header)
            count = 0
            for count, row in enumerate(reader, 1):
                if len(row) != len(headers):
                    raise ValueError(f"{path.name}, row {count + 1}: {len(row)} columns instead of {len(headers)}")
        summaries.append({"source": label, "file": path.name, "rows": count, "columns": headers})
        print(f"{label}: {count} rows, {len(headers)} columns", flush=True)
    return all_headers, summaries


def merge_sources(directory: Path, destination: Path, headers, summaries) -> None:
    total = sum(item["rows"] for item in summaries)
    if total + 1 > 1_048_576:
        raise ValueError(f"{total} data rows exceed one Excel sheet's limit")
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet("Все подписчики")
    sheet.freeze_panes = "A2"
    merged_headers = ["Воронка", *headers]
    for number, header in enumerate(merged_headers, 1):
        sheet.column_dimensions[get_column_letter(number)].width = min(max(len(header) + 3, 16), 32)
    header_cells = []
    for header in merged_headers:
        cell = WriteOnlyCell(sheet, value=header)
        cell.font = Font(bold=True)
        cell.fill = PatternFill(fill_type="solid", fgColor="DCEAF7")
        header_cells.append(cell)
    sheet.append(header_cells)

    for (filename, label, _), summary in zip(SOURCES, summaries, strict=True):
        source, reader = source_reader(directory / f"{filename}.csv")
        with source:
            own_headers = next(reader)
            positions = {name: position for position, name in enumerate(own_headers)}
            for row in reader:
                values = [label, *(row[positions[name]] if name in positions else "" for name in headers)]
                cells = []
                for value in values:
                    cell = WriteOnlyCell(sheet, value=value)
                    cell.data_type = "s"
                    cells.append(cell)
                sheet.append(cells)
        print(f"Added {summary['rows']} rows from {label}", flush=True)
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(merged_headers))}{total + 1}"
    workbook.save(destination)
    print(f"Saved {destination}: {total} rows, {len(merged_headers)} columns", flush=True)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--download", action="store_true")
    parser.add_argument("--download-missing", action="store_true")
    parser.add_argument("--output", type=Path, default=Path("Все_подписчики_воронки.xlsx"))
    args = parser.parse_args()
    csv.field_size_limit(sys.maxsize)
    directory = Path("salebot_sources")
    if args.download or args.download_missing:
        download_sources(directory, missing_only=args.download_missing)
    headers, summaries = inspect_sources(directory)
    merge_sources(directory, args.output, headers, summaries)
    print(json.dumps({"total_rows": sum(x["rows"] for x in summaries), "columns": ["Воронка", *headers], "sources": summaries}, ensure_ascii=False), flush=True)


if __name__ == "__main__":
    main()
