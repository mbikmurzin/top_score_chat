from collections import Counter, defaultdict
from datetime import date, datetime, time
from pathlib import Path

from openpyxl import load_workbook

p = Path(r"C:\Users\User\Downloads\Копия (С1) Выгрузка подписчиков 14.09.xlsx")
w = load_workbook(p, read_only=True, data_only=True)
s = w["Лист2"]
exact = defaultdict(set)
rounded = defaultdict(set)
exact_client = defaultdict(set)
rounded_client = defaultdict(set)
all_ids = defaultdict(list)
winners = {}
for row_number, row in enumerate(s.iter_rows(min_row=2, max_col=11, values_only=False), start=2):
    client = row[1].value
    day = row[8].value.date() if isinstance(row[8].value, datetime) else row[8].value
    clock = row[9].value
    value = row[10].value
    if client is None or not isinstance(day, date):
        continue
    year = 26 if day < date(2026, 5, 1) else 27
    key = (client, year)
    coordinate = f"K{row_number}"
    cell_format = row[10].number_format
    if value is not None:
        all_ids[key].append((coordinate, value, cell_format))
        if row[10].number_format == "0.00E+00":
            rounded[key].add(value)
            rounded_client[client].add(value)
        else:
            exact[key].add(value)
            exact_client[client].add(value)
    moment = datetime.combine(day, clock)
    old = winners.get(key)
    if old is None or moment < old[0]:
        winners[key] = (moment, coordinate, value, cell_format)

counts = Counter()
examples = defaultdict(list)
for key, winner in winners.items():
    direct = winner[2]
    ex = exact[key]
    rnd = rounded[key]
    category = (
        "direct_exact" if direct is not None and winner[3] != "0.00E+00" else
        "direct_rounded" if direct is not None else
        "blank_with_one_exact" if len(ex) == 1 else
        "blank_with_many_exact" if len(ex) > 1 else
        "blank_with_rounded_only" if rnd else
        "blank_no_id"
    )
    counts[category] += 1
    if len(ex) > 1:
        counts["multiple_exact_candidates"] += 1
        if len(examples["multiple_exact_candidates"]) < 10:
            examples["multiple_exact_candidates"].append((key, sorted(ex), all_ids[key][:5]))
    if direct is not None and winner[3] == "0.00E+00" and len(ex) == 1:
        counts["rounded_winner_recoverable"] += 1
    if direct is not None and winner[3] == "0.00E+00" and not ex:
        counts["rounded_winner_unrecoverable"] += 1
    if category == "blank_with_rounded_only":
        counts["blank_rounded_unrecoverable"] += 1
    if len(examples[category]) < 5:
        examples[category].append((key, winner, sorted(ex), sorted(rnd)))

print("winner_count", len(winners))
print("categories", dict(counts))
global_counts = Counter()
global_examples = defaultdict(list)
for key, winner in winners.items():
    candidates = exact_client[key[0]]
    if len(candidates) == 1:
        global_counts["one_exact_for_client"] += 1
    elif len(candidates) > 1:
        global_counts["conflicting_exact_for_client"] += 1
        if len(global_examples["conflicting_exact_for_client"]) < 5:
            global_examples["conflicting_exact_for_client"].append((key, sorted(candidates)))
    elif rounded_client[key[0]]:
        global_counts["rounded_only_for_client"] += 1
    else:
        global_counts["no_id_for_client"] += 1
print("global_categories", dict(global_counts))
print("global_examples", dict(global_examples))
for category, rows in examples.items():
    print(category, rows)
