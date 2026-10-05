"""Keep the earliest row for each (client_id, academic year) in a merged XLSX."""

from collections import Counter
from datetime import date, datetime
from pathlib import Path
import argparse

from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


def academic_year(day: date) -> int | None:
    if date(2026, 3, 1) <= day < date(2026, 4, 1):
        return 26
    if date(2026, 4, 1) <= day < date(2027, 5, 1):
        return 27
    return None


parser = argparse.ArgumentParser()
parser.add_argument("--source", type=Path, default=Path("Все_подписчики_воронки.xlsx"))
parser.add_argument("--output", type=Path, default=Path("Все_подписчики_без_дублей_по_учебному_году.xlsx"))
args = parser.parse_args()

source_book = load_workbook(args.source, read_only=True, data_only=True)
source_sheet = source_book.active
source_rows = source_sheet.iter_rows(values_only=True)
headers = tuple(next(source_rows))
client_index = headers.index("client_id")
date_index = headers.index("current_date")
time_index = headers.index("current_time")

winners = {}
input_count = 0
excluded = Counter()

for source_row_number, row in enumerate(source_rows, start=2):
    input_count += 1
    client_id = row[client_index]
    if client_id is None or str(client_id).strip() == "":
        excluded["empty_client_id"] += 1
        continue
    day = datetime.strptime(row[date_index], "%d.%m.%Y").date()
    year = academic_year(day)
    if year is None:
        excluded["outside_academic_years"] += 1
        continue
    moment = datetime.combine(day, datetime.strptime(row[time_index], "%H:%M").time())
    key = (str(client_id), year)
    previous = winners.get(key)
    if previous is None or moment < previous[0]:
        winners[key] = (moment, source_row_number, row)

selected = sorted(winners.items(), key=lambda item: item[1][1])
year_counts = Counter(key[1] for key, _ in selected)

output_book = Workbook(write_only=True)
output_sheet = output_book.create_sheet("Подписчики")
output_sheet.freeze_panes = "A2"
for index, header in enumerate(headers, start=1):
    output_sheet.column_dimensions[get_column_letter(index)].width = min(max(len(header) + 3, 16), 32)

header_cells = []
for header in headers:
    cell = WriteOnlyCell(output_sheet, value=header)
    cell.font = Font(bold=True)
    cell.fill = PatternFill(fill_type="solid", fgColor="DCEAF7")
    header_cells.append(cell)
output_sheet.append(header_cells)

for _, (_, _, row) in selected:
    cells = []
    for value in row:
        cell = WriteOnlyCell(output_sheet, value=value if value is not None else "")
        cell.data_type = "s"
        cells.append(cell)
    output_sheet.append(cells)

output_sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(selected) + 1}"
output_book.save(args.output)
source_book.close()

print(f"source_rows={input_count}")
print(f"result_rows={len(selected)}")
print(f"removed_duplicates={input_count - sum(excluded.values()) - len(selected)}")
print(f"year_26={year_counts[26]} year_27={year_counts[27]}")
print(f"excluded={dict(excluded)}")
print(f"output={args.output.resolve()}")
