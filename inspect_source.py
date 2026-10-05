from collections import Counter
from pathlib import Path
from zipfile import ZipFile
import re

from openpyxl import load_workbook

p = Path(r"C:\Users\User\Downloads\Копия (С1) Выгрузка подписчиков 14.09.xlsx")
wb = load_workbook(p, read_only=True, data_only=False)
ws = wb["Лист2"]
hits = []
for row in ws.iter_rows(min_row=2, max_col=11, values_only=False):
    cell = row[10]
    if cell.value is not None:
        hits.append((cell.coordinate, cell.value, cell.number_format))
    if len(hits) >= 50:
        break
print("First 50 nonblank:", hits)

with ZipFile(p) as archive:
    xml = archive.read("xl/worksheets/sheet1.xml").decode("utf-8")
for coordinate, value, fmt in hits:
    if fmt == "0.00E+00":
        match = re.search(r'<c\b[^>]*\br="' + coordinate + r'"[^>]*>.*?</c>', xml)
        print("Raw numeric cell:", coordinate, value, match.group(0) if match else "not found")
        if coordinate != hits[-1][0]:
            break
