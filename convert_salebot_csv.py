"""Convert a SaleBot CSV export to XLSX without rounding identifier columns."""

import csv
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


def main(source: Path, destination: Path) -> None:
    csv.field_size_limit(sys.maxsize)
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet("Данные")
    sheet.freeze_panes = "A2"

    with source.open("r", encoding="utf-8-sig", newline="") as input_file:
        reader = csv.reader(input_file, delimiter=";")
        headers = next(reader)
        for column_number, name in enumerate(headers, 1):
            sheet.column_dimensions[get_column_letter(column_number)].width = min(
                max(len(name) + 3, 16), 32
            )

        header_cells = []
        for name in headers:
            cell = WriteOnlyCell(sheet, value=name)
            cell.font = Font(bold=True)
            cell.fill = PatternFill(fill_type="solid", fgColor="DCEAF7")
            header_cells.append(cell)
        sheet.append(header_cells)

        row_count = 0
        for row_count, row in enumerate(reader, 1):
            if len(row) != len(headers):
                raise ValueError(f"Row {row_count + 1} has {len(row)} columns")
            cells = []
            for value in row:
                cell = WriteOnlyCell(sheet, value=value)
                cell.data_type = "s"
                cells.append(cell)
            sheet.append(cells)

    sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{row_count + 1}"
    workbook.save(destination)
    print(f"rows={row_count} columns={len(headers)} output={destination}")


if __name__ == "__main__":
    main(Path(sys.argv[1]), Path(sys.argv[2]))
